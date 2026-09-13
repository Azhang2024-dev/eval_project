"""Canonical Excel styling constants and helpers for kb-qa-eval workbooks."""

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def auto_width(ws, min_width: int = 10, max_width: int = 80) -> None:
    """Set column widths based on cell content, respecting min/max bounds."""
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        best = min_width
        for cell in col_cells:
            if cell.value:
                for line in str(cell.value).split("\n"):
                    best = max(best, min(len(line) + 2, max_width))
        ws.column_dimensions[col_letter].width = best