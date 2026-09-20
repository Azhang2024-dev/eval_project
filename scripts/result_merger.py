"""Normalize and merge evaluation results from multiple models."""

from __future__ import annotations

import hashlib
import json
import math
import os
import re
import statistics
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from scripts.api_client import (
    PROTOCOL_CHAT_COMPLETIONS,
    RESPONSE_MODE_TEXT,
    validate_protocol,
    validate_response_mode,
)
from scripts.eval_stats import (
    BADCASE_THRESHOLD,
    LANGUAGE_META,
    WEIGHTS,
    parse_score,
    score_headers,
    summarize,
    valid_score_count,
    weighted_from_score_dict,
)
from scripts.excel_style import HEADER_FILL, HEADER_FONT, HEADER_ALIGN, auto_width as excel_auto_width

_SECTION_FONT = Font(bold=True, size=12, color="1F4E79")
_ERROR_FONT = Font(color="C00000")


class MergeValidationError(ValueError):
    def __init__(self, message: str, issues: list[dict[str, Any]] | None = None):
        super().__init__(message)
        self.issues = issues or []


@dataclass
class EvaluationRecord:
    query_id: int
    source_df_index: int
    source_excel_row: int
    query: str
    filename: str
    retrieval_status: str
    answer: str
    source_values: dict[str, Any]
    scores: dict[str, int | None]
    reasons: dict[str, str]
    weighted_average: float | None
    overall_summary: str
    status: str
    raw_result: dict[str, Any] | None = None


@dataclass
class ModelArtifact:
    model: str
    language: str
    source_path: str | None
    protocol: str = PROTOCOL_CHAT_COMPLETIONS
    response_mode: str = RESPONSE_MODE_TEXT
    output_path: str | None = None
    checkpoint_path: str | None = None
    input_fingerprint: str = ""
    input_row_count: int = 0
    query_count: int = 0
    base_columns: list[str] = field(default_factory=list)
    source_columns: dict[str, str] = field(default_factory=dict)
    records: dict[int, EvaluationRecord] = field(default_factory=dict)
    metadata_status: str = "metadata_missing"
    issues: list[dict[str, Any]] = field(default_factory=list)
    review_issues: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class MergedResults:
    language: str
    base_columns: list[str]
    input_fingerprint: str
    source_row_count: int
    query_count: int
    records: list[EvaluationRecord]
    artifacts: list[ModelArtifact]
    validation_issues: list[dict[str, Any]] = field(default_factory=list)


def _config(language: str) -> dict[str, Any]:
    meta = LANGUAGE_META.get(language, LANGUAGE_META["en"])
    return {
        "dimensions": meta.dims,
        "score_headers": score_headers(language),
        "summary": meta.summary_key,
    }


def _is_missing(value: Any) -> bool:
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    return isinstance(value, str) and not value.strip()


def _display_value(value: Any) -> str:
    if _is_missing(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _canonical_value(value: Any) -> Any:
    if _is_missing(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    if isinstance(value, (int, str, bool)):
        return value
    if isinstance(value, float):
        return None if math.isnan(value) or math.isinf(value) else round(value, 12)
    return str(value)


def compute_input_fingerprint(df: pd.DataFrame, base_columns: list[str] | None = None) -> str:
    columns = list(base_columns or df.columns)
    payload = {
        "columns": columns,
        "rows": [[_canonical_value(value) for value in row] for row in df[columns].itertuples(index=False, name=None)],
    }
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _find_column(columns: list[str], patterns: tuple[str, ...]) -> str | None:
    for column in columns:
        text = str(column).strip()
        for pattern in patterns:
            if pattern.casefold() == text.casefold() or re.search(pattern, text, re.IGNORECASE):
                return column
    return None


def _source_columns(base_columns: list[str]) -> dict[str, str]:
    return {
        "query": _find_column(base_columns, ("query", "question", "问题")) or "",
        "filename": _find_column(base_columns, ("filename", "文件名", "source_file", "来源")) or "",
        "answer": _find_column(base_columns, ("answer", "response", "回答", "大模型问答结果")) or "",
        "chunks": _find_column(base_columns, ("搜索结果", "retrieved", "context", "chunks")) or "",
        "reference": _find_column(base_columns, ("标注", "reference", "gold", "标准答案")) or "",
        "retrieval_score": _find_column(base_columns, ("score", "相似度", "similarity")) or "",
    }


def _retrieval_status(row: dict[str, Any], columns: dict[str, str]) -> str:
    filename = row.get(columns.get("filename", "")) if columns.get("filename") else None
    chunks = row.get(columns.get("chunks", "")) if columns.get("chunks") else None
    filename_missing = _is_missing(filename)
    chunks_missing = _is_missing(chunks)
    if filename_missing and chunks_missing:
        return "missing_filename_and_search_result"
    if filename_missing:
        return "missing_filename"
    if chunks_missing:
        return "missing_search_result"
    return "available"





def _infer_status(answer: Any, scores: dict[str, int | None], raw_status: str | None = None) -> str:
    if raw_status in {"success", "empty", "failed", "not_evaluated", "incomplete_result"}:
        return raw_status
    count = valid_score_count(scores)
    if count == 5:
        return "empty" if _is_missing(answer) else "success"
    return "not_evaluated" if count == 0 else "incomplete_result"


def _infer_model_label(path: str) -> str:
    stem = Path(path).stem
    label = stem.split("_evaluated_", 1)[1] if "_evaluated_" in stem else stem
    for suffix in ("_anthropic_messages", "_chat_completions", "_responses"):
        if label.endswith(suffix):
            label = label[: -len(suffix)]
    return label or stem


def infer_protocol(path: str) -> str:
    stem = Path(path).stem.casefold()
    if stem.endswith("_anthropic_messages"):
        return "anthropic_messages"
    if stem.endswith("_responses"):
        return "responses"
    return PROTOCOL_CHAT_COMPLETIONS


def _detect_language(headers: list[str]) -> str:
    for language, meta in LANGUAGE_META.items():
        headers_for_lang = score_headers(language)
        if any(score_header in headers for score_header, _ in headers_for_lang):
            return language
    raise MergeValidationError("Could not detect evaluation language from score headers", [{"type": "unknown_language"}])


def _score_header_map(headers: list[str], language: str) -> tuple[list[str], dict[str, tuple[str, str]]]:
    config = _config(language)
    missing = [score for score, _ in config["score_headers"] if score not in headers]
    if missing:
        raise MergeValidationError(
            f"Missing score columns for {language}: {', '.join(missing)}",
            [{"type": "missing_score_columns", "columns": missing}],
        )
    return [score for score, _ in config["score_headers"]], {
        dimension: pair for dimension, pair in zip(config["dimensions"], config["score_headers"])
    }


def _base_columns_for_headers(headers: list[str], language: str) -> list[str]:
    score_headers, _ = _score_header_map(headers, language)
    first_score_index = min(headers.index(header) for header in score_headers)
    base_columns = headers[:first_score_index]
    if not _source_columns(base_columns).get("query"):
        raise MergeValidationError("Evaluation workbook does not contain a query column", [{"type": "missing_query_column"}])
    return base_columns


def _read_sheet(path: str) -> tuple[str, pd.DataFrame]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        candidates = {name for name in ("Evaluation", "评测结果", "Évaluation") if name in wb.sheetnames}
        sheet_name = next(iter(candidates), wb.sheetnames[0])
    finally:
        wb.close()
    return sheet_name, pd.read_excel(path, sheet_name=sheet_name)


def _record_from_row(row: dict[str, Any], source_df_index: int, query_id: int, language: str, score_mapping: dict[str, tuple[str, str]], source_columns: dict[str, str], raw_status: str | None = None, raw_result: dict[str, Any] | None = None) -> EvaluationRecord:
    config = _config(language)
    scores = {}
    reasons = {}
    for dimension in config["dimensions"]:
        score_header, reason_header = score_mapping[dimension]
        scores[dimension] = parse_score(row.get(score_header))
        reasons[dimension] = _display_value(row.get(reason_header))
    summary_header = config["summary"]
    provided_weighted = row.get("Weighted Average", row.get("加权均分"))
    computed_weighted = weighted_from_score_dict(scores, language)
    weighted = computed_weighted
    if weighted is None and not _is_missing(provided_weighted):
        try:
            weighted = round(float(provided_weighted), 2)
        except (TypeError, ValueError):
            weighted = None
    answer = row.get(source_columns["answer"]) if source_columns["answer"] else None
    return EvaluationRecord(
        query_id=query_id,
        source_df_index=source_df_index,
        source_excel_row=source_df_index + 2,
        query=_display_value(row.get(source_columns["query"])) if source_columns["query"] else "",
        filename=_display_value(row.get(source_columns["filename"])) if source_columns["filename"] else "",
        retrieval_status=_retrieval_status(row, source_columns),
        answer=_display_value(answer),
        source_values=row,
        scores=scores,
        reasons=reasons,
        weighted_average=weighted,
        overall_summary=_display_value(row.get(summary_header)),
        status=_infer_status(answer, scores, raw_status),
        raw_result=raw_result,
    )


def _record_from_runtime(source_row: dict[str, Any], source_df_index: int, query_id: int, item: dict[str, Any] | None, language: str, source_columns: dict[str, str]) -> EvaluationRecord:
    config = _config(language)
    result = item.get("result") if item and isinstance(item.get("result"), dict) else {}
    scores = {}
    reasons = {}
    for dimension in config["dimensions"]:
        dim_data = result.get(dimension, {})
        dim_data = dim_data if isinstance(dim_data, dict) else {}
        scores[dimension] = parse_score(dim_data.get("score" if language == "en" else ("分数" if language == "zh" else "note")))
        reasons[dimension] = _display_value(dim_data.get("reason" if language == "en" else ("理由" if language == "zh" else "raison")))
    # evaluate_batch result items only carry index/result/status — query,
    # filename and answer must fall back to the source row so runtime-built
    # records keep their identity (previously query became '' and every record
    # after the first was flagged as a spurious duplicate_query_text).
    query = (item or {}).get("query") or source_row.get(source_columns["query"])
    filename = (item or {}).get("filename") or source_row.get(source_columns["filename"])
    answer = (item or {}).get("answer") or source_row.get(source_columns["answer"])
    status = item.get("status") if item else "not_evaluated"
    return EvaluationRecord(
        query_id=query_id,
        source_df_index=source_df_index,
        source_excel_row=source_df_index + 2,
        query=_display_value(query),
        filename=_display_value(filename),
        retrieval_status=_retrieval_status(source_row, source_columns),
        answer=_display_value(answer),
        source_values=source_row,
        scores=scores,
        reasons=reasons,
        weighted_average=weighted_from_score_dict(scores, language),
        overall_summary=_display_value(result.get(config["summary"], "")),
        status=_infer_status(answer, scores, status),
        raw_result=result or None,
    )


def _add_record_issues(artifact: ModelArtifact) -> None:
    seen = {}
    for record in artifact.records.values():
        if record.query in seen:
            artifact.issues.append({"type": "duplicate_query_text", "model": artifact.model, "source_df_index": record.source_df_index, "other_source_df_index": seen[record.query], "query": record.query})
        else:
            seen[record.query] = record.source_df_index
        if record.status == "incomplete_result":
            artifact.issues.append({"type": "incomplete_score_result", "model": artifact.model, "source_df_index": record.source_df_index, "source_excel_row": record.source_excel_row, "valid_score_count": valid_score_count(record.scores)})


def _read_review_consistency(path: str, artifact: ModelArtifact) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        review_sheet_name = LANGUAGE_META.get(artifact.language, LANGUAGE_META["en"]).review_sheet_name
        if review_sheet_name not in wb.sheetnames:
            return []
        ws = wb[review_sheet_name]
        headers = [ws.cell(1, column).value for column in range(1, ws.max_column + 1)]
        h = {str(value): index + 1 for index, value in enumerate(headers) if value is not None}
        required = ["行号", "RR 原分", "CA 原分", "SC 原分", "IF 原分", "RN 原分"]
        if any(key not in h for key in required):
            return []
        dimensions = _config(artifact.language)["dimensions"]
        issues = []
        for row_number in range(2, ws.max_row + 1):
            try:
                source_index = int(ws.cell(row_number, h["行号"]).value)
            except (TypeError, ValueError):
                continue
            record = artifact.records.get(source_index)
            if record is None:
                issues.append({"type": "review_unknown_source_index", "model": artifact.model, "review_sheet_row": row_number, "source_df_index": source_index})
                continue
            review_scores = [parse_score(ws.cell(row_number, h[key]).value) for key in ("RR 原分", "CA 原分", "SC 原分", "IF 原分", "RN 原分")]
            evaluation_scores = [record.scores.get(dimension) for dimension in dimensions]
            if review_scores != evaluation_scores:
                issues.append({"type": "review_score_mismatch", "model": artifact.model, "review_sheet_row": row_number, "source_df_index": source_index, "source_excel_row": source_index + 2, "review_scores": review_scores, "evaluation_scores": evaluation_scores})
        return issues
    finally:
        wb.close()


def load_evaluation_artifact(path: str, model: str | None = None, protocol: str | None = None, response_mode: str | None = None, expected_language: str | None = None) -> ModelArtifact:
    path = os.path.abspath(path)
    if not os.path.isfile(path):
        raise FileNotFoundError(path)
    _, df = _read_sheet(path)
    headers = [str(column) for column in df.columns]
    language = _detect_language(headers)
    if expected_language and language != expected_language:
        raise MergeValidationError(f"Language mismatch for {path}: {language} != {expected_language}", [{"type": "language_mismatch", "path": path, "detected": language, "expected": expected_language}])
    base_columns = _base_columns_for_headers(headers, language)
    source_columns = _source_columns(base_columns)
    _, score_mapping = _score_header_map(headers, language)
    model = model or _infer_model_label(path)
    protocol = validate_protocol(protocol or infer_protocol(path))
    response_mode = validate_response_mode(response_mode or RESPONSE_MODE_TEXT)
    artifact = ModelArtifact(model=model, language=language, source_path=path, protocol=protocol, response_mode=response_mode, output_path=path, input_fingerprint=compute_input_fingerprint(df, base_columns), input_row_count=len(df), base_columns=base_columns, source_columns=source_columns)
    artifact.issues.append({"type": "metadata_missing", "model": model, "path": path})
    query_indices = [index for index, value in enumerate(df[source_columns["query"]].tolist()) if not _is_missing(value)]
    artifact.query_count = len(query_indices)
    for query_id, source_index in enumerate(query_indices, start=1):
        row = {str(key): value for key, value in df.iloc[source_index].to_dict().items()}
        artifact.records[source_index] = _record_from_row(row, source_index, query_id, language, score_mapping, source_columns)
        provided = row.get("Weighted Average", row.get("加权均分"))
        computed = artifact.records[source_index].weighted_average
        if computed is not None and not _is_missing(provided):
            try:
                if abs(float(provided) - computed) > 0.02:
                    artifact.issues.append({"type": "weighted_average_mismatch", "model": model, "source_df_index": source_index, "source_excel_row": source_index + 2, "provided": provided, "computed": computed})
            except (TypeError, ValueError):
                artifact.issues.append({"type": "invalid_weighted_average", "model": model, "source_df_index": source_index, "source_excel_row": source_index + 2, "provided": provided})
    _add_record_issues(artifact)
    artifact.review_issues = _read_review_consistency(path, artifact)
    return artifact


def build_runtime_artifact(model: str, source_df: pd.DataFrame, columns: dict[str, Any], eval_data: list[dict[str, Any]], results: list[dict[str, Any]], language: str, protocol: str = PROTOCOL_CHAT_COMPLETIONS, response_mode: str = RESPONSE_MODE_TEXT, source_path: str | None = None, output_path: str | None = None, checkpoint_path: str | None = None) -> ModelArtifact:
    protocol = validate_protocol(protocol)
    response_mode = validate_response_mode(response_mode)
    base_columns = [str(column) for column in source_df.columns]
    source_columns = {key: str(value or "") if str(value or "") in base_columns else "" for key, value in {
        "query": columns.get("query"), "filename": columns.get("filename"), "answer": columns.get("answer"), "chunks": columns.get("chunks"), "reference": columns.get("reference"), "retrieval_score": columns.get("score")
    }.items()}
    result_map = {int(item["index"]): item for item in results if "index" in item}
    eval_map = {int(item["index"]): item for item in eval_data if "index" in item}
    query_indices = [index for index, value in enumerate(source_df[source_columns["query"]].tolist()) if not _is_missing(value)]
    artifact = ModelArtifact(model=model, language=language, source_path=os.path.abspath(source_path) if source_path else None, protocol=protocol, response_mode=response_mode, output_path=os.path.abspath(output_path) if output_path else None, checkpoint_path=os.path.abspath(checkpoint_path) if checkpoint_path else None, input_fingerprint=compute_input_fingerprint(source_df, base_columns), input_row_count=len(source_df), query_count=len(query_indices), base_columns=base_columns, source_columns=source_columns, metadata_status="runtime_metadata")
    for query_id, source_index in enumerate(query_indices, start=1):
        source_row = {str(key): value for key, value in source_df.iloc[source_index].to_dict().items()}
        item = result_map.get(source_index, eval_map.get(source_index))
        artifact.records[source_index] = _record_from_runtime(source_row, source_index, query_id, item, language, source_columns)
    if checkpoint_path:
        artifact.issues.append({"type": "checkpoint_path", "model": model, "path": os.path.abspath(checkpoint_path)})
    _add_record_issues(artifact)
    return artifact


def _compare_source_identity(reference: ModelArtifact, candidate: ModelArtifact) -> list[dict[str, Any]]:
    issues = []
    if reference.input_fingerprint != candidate.input_fingerprint:
        issues.append({"type": "input_fingerprint_mismatch", "reference_model": reference.model, "candidate_model": candidate.model, "reference": reference.input_fingerprint, "candidate": candidate.input_fingerprint})
    if reference.base_columns != candidate.base_columns:
        issues.append({"type": "base_columns_mismatch", "reference_model": reference.model, "candidate_model": candidate.model})
    left_indices = set(reference.records)
    right_indices = set(candidate.records)
    for source_index in sorted(left_indices - right_indices):
        issues.append({"type": "missing_query", "model": candidate.model, "source_df_index": source_index})
    for source_index in sorted(right_indices - left_indices):
        issues.append({"type": "extra_query", "model": candidate.model, "source_df_index": source_index})
    for source_index in sorted(left_indices & right_indices):
        left = reference.records[source_index]
        right = candidate.records[source_index]
        if left.query != right.query:
            issues.append({"type": "query_mismatch", "model": candidate.model, "source_df_index": source_index, "reference_query": left.query, "candidate_query": right.query})
        for field_name in ("filename", "answer", "retrieval_status"):
            if getattr(left, field_name) != getattr(right, field_name):
                issues.append({"type": "source_field_mismatch", "model": candidate.model, "field": field_name, "source_df_index": source_index, "reference": getattr(left, field_name), "candidate": getattr(right, field_name)})
    return issues


def merge_artifacts(artifacts: list[ModelArtifact]) -> MergedResults:
    if not artifacts:
        raise MergeValidationError("At least one model artifact is required")
    model_names = [artifact.model for artifact in artifacts]
    duplicates = sorted({model for model in model_names if model_names.count(model) > 1})
    if duplicates:
        raise MergeValidationError(f"Duplicate model labels: {', '.join(duplicates)}", [{"type": "duplicate_model", "model": model} for model in duplicates])
    languages = {artifact.language for artifact in artifacts}
    if len(languages) != 1:
        raise MergeValidationError("All model artifacts must use the same language", [{"type": "language_mismatch", "languages": sorted(languages)}])
    reference = artifacts[0]
    issues = []
    for artifact in artifacts:
        issues.extend(artifact.issues)
        issues.extend(artifact.review_issues)
    for artifact in artifacts[1:]:
        issues.extend(_compare_source_identity(reference, artifact))
    fatal_types = {"input_fingerprint_mismatch", "base_columns_mismatch", "missing_query", "extra_query", "query_mismatch", "source_field_mismatch", "language_mismatch", "duplicate_model"}
    if any(issue.get("type") in fatal_types for issue in issues):
        raise MergeValidationError("Model artifacts do not describe the same input dataset; merge was refused", issues)
    return MergedResults(language=reference.language, base_columns=reference.base_columns, input_fingerprint=reference.input_fingerprint, source_row_count=reference.input_row_count, query_count=reference.query_count, records=[reference.records[index] for index in sorted(reference.records)], artifacts=artifacts, validation_issues=issues)


def _format_number(value: Any) -> Any:
    if value is None or _is_missing(value):
        return ""
    return round(value, 4) if isinstance(value, float) else value





def _write_header(ws, row: int, column: int, value: Any) -> None:
    cell = ws.cell(row=row, column=column, value=value)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = HEADER_ALIGN


def _model_label(model: str, dimension: str) -> str:
    return f"{model}-{dimension}"


def _base_value(record: EvaluationRecord, field_name: str) -> Any:
    if field_name == "query_id":
        return record.query_id
    if field_name == "source_df_index":
        return record.source_df_index
    if field_name == "source_excel_row":
        return record.source_excel_row
    if field_name == "query":
        return record.query
    if field_name == "filename":
        return record.filename
    if field_name == "retrieval_status":
        return record.retrieval_status
    if field_name == "answer":
        return record.answer
    return record.source_values.get(field_name)


def write_merged_workbook(merged: MergedResults, output_path: str) -> str:
    output_path = os.path.abspath(output_path)
    source_paths = {os.path.abspath(path) for artifact in merged.artifacts for path in (artifact.source_path, artifact.output_path) if path}
    if output_path in source_paths:
        raise MergeValidationError("Merged output must not overwrite a source workbook", [{"type": "output_overwrites_source", "path": output_path}])
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "逐条模型对比"
    config = _config(merged.language)
    dimensions = config["dimensions"]
    score_headers = config["score_headers"]
    summary_header = config["summary"]

    # Keep the original input columns first, then place every model's complete
    # score/reason block side by side in the same table.
    base_columns = list(merged.base_columns)
    if not base_columns and merged.records:
        base_columns = list(merged.records[0].source_values)
    headers = list(base_columns)
    for artifact in merged.artifacts:
        for score_header, reason_header in score_headers:
            headers.extend([
                f"{artifact.model}-{score_header}",
                f"{artifact.model}-{reason_header}",
            ])
        headers.extend([
            f"{artifact.model}-Weighted Average",
            f"{artifact.model}-{summary_header}",
        ])

    for column, header in enumerate(headers, start=1):
        _write_header(ws, 1, column, header)

    by_model = {artifact.model: artifact.records for artifact in merged.artifacts}
    for row_number, reference_record in enumerate(merged.records, start=2):
        column = 1
        for base_column in base_columns:
            value = reference_record.source_values.get(base_column)
            ws.cell(row=row_number, column=column, value="" if _is_missing(value) else value)
            column += 1

        for artifact in merged.artifacts:
            record = by_model[artifact.model].get(reference_record.source_df_index)
            if record is None:
                record = EvaluationRecord(
                    reference_record.query_id,
                    reference_record.source_df_index,
                    reference_record.source_excel_row,
                    reference_record.query,
                    reference_record.filename,
                    reference_record.retrieval_status,
                    reference_record.answer,
                    reference_record.source_values,
                    {dimension: None for dimension in dimensions},
                    {dimension: "" for dimension in dimensions},
                    None,
                    "",
                    "missing_record",
                )

            for dimension, (score_header, reason_header) in zip(dimensions, score_headers):
                ws.cell(row=row_number, column=column, value=record.scores.get(dimension))
                column += 1
                ws.cell(row=row_number, column=column, value=record.reasons.get(dimension, ""))
                column += 1
            ws.cell(row=row_number, column=column, value=_format_number(record.weighted_average))
            column += 1
            ws.cell(row=row_number, column=column, value=record.overall_summary or "")
            column += 1

    excel_auto_width(ws, min_width=12, max_width=70)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    stats_sheet = wb.create_sheet("模型统计")
    stats_sheet.cell(row=1, column=1, value="多模型评测统计").font = _SECTION_FONT
    for row, (label, value) in enumerate([("输入指纹", merged.input_fingerprint), ("原始行数", merged.source_row_count), ("Query数量", merged.query_count)], start=2):
        stats_sheet.cell(row=row, column=1, value=label).font = Font(bold=True)
        stats_sheet.cell(row=row, column=2, value=value)
    stats_row = 6
    stats_headers = ["模型", "协议", "语言", "metadata状态", "评测总数", "成功", "回答为空", "失败/未完成"]
    for dimension in dimensions:
        stats_headers.extend([f"{dimension}-均值", f"{dimension}-中位数", f"{dimension}-标准差", f"{dimension}-最低", f"{dimension}-最高"])
    stats_headers.extend(["Weighted Average-均值(全部行失败计0)", "Weighted Average-均值(仅计分)", "Weighted Average-最低(仅计分)", "Weighted Average-最高(仅计分)", "Badcase数量", "Badcase率(仅计分母)", "failed_penalty_policy"])
    for column, header in enumerate(stats_headers, start=1):
        _write_header(stats_sheet, stats_row, column, header)
    stats_row += 1
    for artifact in merged.artifacts:
        records = list(artifact.records.values())
        valid_records = [record for record in records if record.status in {"success", "empty"}]
        weighted = [float(record.weighted_average) for record in valid_records if record.weighted_average is not None]
        values = [artifact.model, artifact.protocol, artifact.language, artifact.metadata_status, len(records), sum(record.status == "success" for record in records), sum(record.status == "empty" for record in records), sum(record.status not in {"success", "empty"} for record in records)]
        for dimension in dimensions:
            dimension_values = [record.scores[dimension] for record in valid_records if record.scores.get(dimension) is not None]
            dim_stats = summarize([float(v) for v in dimension_values])
            mean_value = round(dim_stats["mean"], 4)
            min_value = round(dim_stats["min"], 4)
            max_value = round(dim_stats["max"], 4)
            std_value = round(dim_stats["stdev"], 4)
            median_value = round(statistics.median(dimension_values), 4) if dimension_values else ""
            values.extend([mean_value, median_value, std_value, min_value, max_value])
        badcases = sum(value < BADCASE_THRESHOLD for value in weighted)
        # P2.5 honesty: headline mean treats failed rows as 0; scored-only mean labeled separately.
        n_failed = sum(record.status not in {"success", "empty"} for record in records)
        all_row_scores = list(weighted) + [0.0] * n_failed
        mean_all = round(statistics.mean(all_row_scores), 4) if all_row_scores else ""
        mean_scored = round(statistics.mean(weighted), 4) if weighted else ""
        policy = "failed/non-scored rows contribute 0.0 to all-rows headline mean; scored-only mean excludes them"
        values.extend([
            mean_all,
            mean_scored,
            min(weighted) if weighted else "",
            max(weighted) if weighted else "",
            badcases,
            f"{badcases / len(weighted) * 100:.1f}%" if weighted else "0.0%",
            policy,
        ])
        for column, value in enumerate(values, start=1):
            stats_sheet.cell(row=stats_row, column=column, value=value)
        stats_row += 1
    excel_auto_width(stats_sheet, min_width=12, max_width=70)
    stats_sheet.freeze_panes = "A7"
    stats_sheet.auto_filter.ref = stats_sheet.dimensions

    validation = wb.create_sheet("数据校验")
    validation.cell(row=1, column=1, value="合并数据校验").font = _SECTION_FONT
    for row, (label, value) in enumerate([("输入指纹", merged.input_fingerprint), ("语言", merged.language), ("模型数量", len(merged.artifacts)), ("原始行数", merged.source_row_count), ("Query数量", merged.query_count)], start=3):
        validation.cell(row=row, column=1, value=label).font = Font(bold=True)
        validation.cell(row=row, column=2, value=value)
    row = 10
    artifact_headers = ["模型", "协议", "响应模式", "语言", "metadata状态", "输入指纹", "原始行数", "Query数量", "源文件", "输出文件", "checkpoint"]
    for column, header in enumerate(artifact_headers, start=1):
        _write_header(validation, row, column, header)
    row += 1
    for artifact in merged.artifacts:
        values = [artifact.model, artifact.protocol, artifact.response_mode, artifact.language, artifact.metadata_status, artifact.input_fingerprint, artifact.input_row_count, artifact.query_count, artifact.source_path or "", artifact.output_path or "", artifact.checkpoint_path or ""]
        for column, value in enumerate(values, start=1):
            validation.cell(row=row, column=column, value=value)
        row += 1
    row += 1
    validation_headers = ["级别", "类型", "模型", "原始DataFrame索引", "原始Excel行号", "详情"]
    for column, header in enumerate(validation_headers, start=1):
        _write_header(validation, row, column, header)
    row += 1
    warning_types = {"metadata_missing", "weighted_average_mismatch", "review_score_mismatch", "incomplete_score_result", "duplicate_query_text", "checkpoint_path"}
    if not merged.validation_issues:
        validation.cell(row=row, column=1, value="INFO")
        validation.cell(row=row, column=2, value="no_issues")
    else:
        for issue in merged.validation_issues:
            issue_type = issue.get("type", "unknown")
            level = "WARNING" if issue_type in warning_types else "ERROR"
            source_index = issue.get("source_df_index", "")
            excel_row = issue.get("source_excel_row", source_index + 2 if isinstance(source_index, int) else "")
            values = [level, issue_type, issue.get("model", ""), source_index, excel_row, json.dumps(issue, ensure_ascii=False, sort_keys=True)]
            for column, value in enumerate(values, start=1):
                cell = validation.cell(row=row, column=column, value=value)
                if level == "ERROR":
                    cell.font = _ERROR_FONT
            row += 1
    excel_auto_width(validation, min_width=12, max_width=100)
    validation.freeze_panes = "A11"
    validation.auto_filter.ref = validation.dimensions

    wb.save(output_path)
    return output_path


def merge_evaluation_files(files: list[tuple[str, str]], output_path: str, protocols: dict[str, str] | None = None, response_modes: dict[str, str] | None = None, language: str | None = None) -> str:
    protocols = protocols or {}
    response_modes = response_modes or {}
    artifacts = [load_evaluation_artifact(path, model=model, protocol=protocols.get(model), response_mode=response_modes.get(model), expected_language=language) for model, path in files]
    return write_merged_workbook(merge_artifacts(artifacts), output_path)
