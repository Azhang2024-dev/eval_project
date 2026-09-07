from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from scripts.result_merger import (
    MergeValidationError,
    build_runtime_artifact,
    load_evaluation_artifact,
    merge_artifacts,
    write_merged_workbook,
)


BASE_HEADERS = ["query", "filename", "搜索结果", "score", "大模型问答结果", "标注"]
SCORE_PAIRS = [
    ("Response Relevance-Score", "Response Relevance-Reason"),
    ("Citation Accuracy-Score", "Citation Accuracy-Reason"),
    ("Scope Completeness-Score", "Scope Completeness-Reason"),
    ("Instruction Following-Score", "Instruction Following-Reason"),
    ("Readability and Naturalness-Score", "Readability and Naturalness-Reason"),
]
DIMENSIONS = ["Response Relevance", "Citation Accuracy", "Scope Completeness", "Instruction Following", "Readability and Naturalness"]


def _write_workbook(path: str, rows: list[dict], scores: dict[int, list[int | None]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Evaluation"
    headers = BASE_HEADERS[:]
    for score_header, reason_header in SCORE_PAIRS:
        headers.extend([score_header, reason_header])
    headers.extend(["Weighted Average", "Overall Summary"])
    ws.append(headers)
    for index, row in enumerate(rows):
        values = [row.get(header) for header in BASE_HEADERS]
        row_scores = scores.get(index, [None] * 5)
        for score, dimension in zip(row_scores, DIMENSIONS):
            values.extend([score, f"reason-{dimension}" if score is not None else ""])
        weighted = sum(score * weight for score, weight in zip(row_scores, [0.3, 0.3, 0.15, 0.15, 0.1])) if all(score is not None for score in row_scores) else None
        values.extend([weighted, "summary"])
        ws.append(values)
    wb.save(path)


class ResultMergerTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.dir = Path(self.temp.name)
        self.rows = [
            {"query": "q0", "filename": "a.md", "搜索结果": "ctx0", "score": 0.9, "大模型问答结果": "a0", "标注": "gold0"},
            {"query": None, "filename": "a.md", "搜索结果": "ctx-continuation", "score": 0.8, "大模型问答结果": None, "标注": None},
            {"query": "q2", "filename": None, "搜索结果": None, "score": 0.0, "大模型问答结果": "a2", "标注": "gold2"},
        ]
        self.scores_a = {0: [5, 4, 5, 5, 4], 2: [1, 1, 1, 1, 1]}
        self.scores_b = {0: [4, 4, 5, 5, 5], 2: [2, 1, 1, 1, 1]}

    def tearDown(self):
        self.temp.cleanup()

    def test_merge_aligns_by_source_index_and_excel_row(self):
        first = self.dir / "input_evaluated_model-a.xlsx"
        second = self.dir / "input_evaluated_model-b.xlsx"
        _write_workbook(str(first), self.rows, self.scores_a)
        _write_workbook(str(second), self.rows, self.scores_b)
        merged = merge_artifacts([
            load_evaluation_artifact(str(first), model="model-a"),
            load_evaluation_artifact(str(second), model="model-b"),
        ])
        self.assertEqual(merged.query_count, 2)
        q2 = next(record for record in merged.records if record.source_df_index == 2)
        self.assertEqual(q2.source_excel_row, 4)
        self.assertEqual(q2.query, "q2")
        self.assertEqual(q2.retrieval_status, "missing_filename_and_search_result")
        self.assertEqual(merged.artifacts[0].records[2].scores[DIMENSIONS[0]], 1)
        self.assertEqual(merged.artifacts[1].records[2].scores[DIMENSIONS[0]], 2)

    def test_output_contains_one_wide_result_sheet_with_scores_and_reasons(self):
        first = self.dir / "input_evaluated_model-a.xlsx"
        second = self.dir / "input_evaluated_model-b.xlsx"
        _write_workbook(str(first), self.rows, self.scores_a)
        _write_workbook(str(second), self.rows, self.scores_b)
        merged = merge_artifacts([
            load_evaluation_artifact(str(first), model="model-a"),
            load_evaluation_artifact(str(second), model="model-b"),
        ])
        output = self.dir / "merged.xlsx"
        write_merged_workbook(merged, str(output))
        wb = load_workbook(output, data_only=True)
        self.assertEqual(wb.sheetnames, ["逐条模型对比", "模型统计", "数据校验"])
        ws = wb["逐条模型对比"]
        headers = [ws.cell(1, column).value for column in range(1, ws.max_column + 1)]
        self.assertEqual(headers[:6], BASE_HEADERS)
        self.assertIn("model-a-Response Relevance-Score", headers)
        self.assertIn("model-a-Response Relevance-Reason", headers)
        self.assertIn("model-a-Weighted Average", headers)
        self.assertIn("model-a-Overall Summary", headers)
        self.assertIn("model-b-Citation Accuracy-Reason", headers)
        self.assertNotIn("Weighted Range", headers)
        row = next(row for row in range(2, ws.max_row + 1) if ws.cell(row, 1).value == "q2")
        self.assertEqual(ws.cell(row, 4).value, 0.0)
        model_a_rr_score = headers.index("model-a-Response Relevance-Score") + 1
        model_a_rr_reason = headers.index("model-a-Response Relevance-Reason") + 1
        self.assertEqual(ws.cell(row, model_a_rr_score).value, 1)
        self.assertEqual(ws.cell(row, model_a_rr_reason).value, "reason-Response Relevance")

    def test_mismatched_input_is_refused(self):
        first = self.dir / "input_evaluated_model-a.xlsx"
        second = self.dir / "different_evaluated_model-b.xlsx"
        _write_workbook(str(first), self.rows, self.scores_a)
        changed = list(self.rows)
        changed[0] = dict(changed[0], query="different")
        _write_workbook(str(second), changed, self.scores_b)
        with self.assertRaises(MergeValidationError) as context:
            merge_artifacts([
                load_evaluation_artifact(str(first), model="model-a"),
                load_evaluation_artifact(str(second), model="model-b"),
            ])
        self.assertTrue(any(issue["type"] == "input_fingerprint_mismatch" for issue in context.exception.issues))

    def test_output_cannot_overwrite_source_workbook(self):
        first = self.dir / "input_evaluated_model-a.xlsx"
        _write_workbook(str(first), self.rows, self.scores_a)
        artifact = load_evaluation_artifact(str(first), model="model-a")
        merged = merge_artifacts([artifact])
        with self.assertRaises(MergeValidationError):
            write_merged_workbook(merged, str(first))

    def test_runtime_artifact_maps_index_812_to_excel_row_814(self):
        rows = [{"query": None, "filename": "f.md", "搜索结果": "ctx", "score": 0.5, "大模型问答结果": "a", "标注": "g"} for _ in range(813)]
        rows[812]["query"] = "LLAMA B"
        source_df = pd.DataFrame(rows)
        columns = {"query": "query", "filename": "filename", "chunks": "搜索结果", "score": "score", "answer": "大模型问答结果", "reference": "标注"}
        result = {"index": 812, "status": "success", "result": {dimension: {"score": 5, "reason": "ok"} for dimension in DIMENSIONS} | {"Overall Summary": "ok"}}
        artifact = build_runtime_artifact("model", source_df, columns, [{"index": 812, "query": "LLAMA B", "filename": "f.md", "answer": "a", "chunks": []}], [result], "en")
        self.assertEqual(artifact.records[812].source_excel_row, 814)

    def test_runtime_artifact_preserves_query_and_answer(self):
        """Regression: evaluate_batch result items lack query/answer — the
        record must fall back to the source row, or every record after the
        first gets a spurious duplicate_query_text issue."""
        rows = [
            {"query": "q0", "filename": "a.md", "搜索结果": "ctx0", "score": 0.9, "大模型问答结果": "a0", "标注": "g0"},
            {"query": "q1", "filename": "a.md", "搜索结果": "ctx1", "score": 0.8, "大模型问答结果": "a1", "标注": "g1"},
        ]
        source_df = pd.DataFrame(rows)
        columns = {"query": "query", "filename": "filename", "chunks": "搜索结果", "score": "score", "answer": "大模型问答结果", "reference": "标注"}
        eval_data = [
            {"index": 0, "query": "q0", "filename": "a.md", "answer": "a0", "chunks": []},
            {"index": 1, "query": "q1", "filename": "a.md", "answer": "a1", "chunks": []},
        ]
        results = [
            {"index": 0, "status": "success", "result": {dimension: {"score": 5, "reason": "ok"} for dimension in DIMENSIONS} | {"Overall Summary": "ok"}},
            {"index": 1, "status": "success", "result": {dimension: {"score": 4, "reason": "ok"} for dimension in DIMENSIONS} | {"Overall Summary": "ok"}},
        ]
        artifact = build_runtime_artifact("model", source_df, columns, eval_data, results, "en")
        self.assertEqual(artifact.records[0].query, "q0")
        self.assertEqual(artifact.records[1].query, "q1")
        self.assertEqual(artifact.records[0].answer, "a0")
        self.assertFalse(any(issue["type"] == "duplicate_query_text" for issue in artifact.issues))


if __name__ == "__main__":
    unittest.main()
